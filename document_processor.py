
import os
import json
import pandas as pd
from typing import List, Dict, Any, Optional
import streamlit as st
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.schema import Document
import logging

# File format specific imports
try:
    from pypdf import PdfReader
except ImportError:
    st.error("PDF processing libraries not installed. Please install pypdf.")

try:
    from docx import Document as DocxDocument
except ImportError:
    st.error("DOCX processing library not installed. Please install python-docx.")

try:
    from pptx import Presentation
except ImportError:
    st.error("PPTX processing library not installed. Please install python-pptx.")

try:
    from bs4 import BeautifulSoup
except ImportError:
    st.error("HTML processing library not installed. Please install beautifulsoup4.")

logger = logging.getLogger(__name__)

class DocumentProcessor:
    """Process various document formats and convert to Langchain Documents"""
    
    def __init__(self, chunk_size: int = 1000, chunk_overlap: int = 200):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            length_function=len,
            separators=["\n\n", "\n", " ", ""]
        )
    
    def process_file(self, uploaded_file) -> List[Document]:
        """Process uploaded file based on its format"""
        try:
            file_extension = uploaded_file.name.split('.')[-1].lower()
            
            # Create metadata
            metadata = {
                "filename": uploaded_file.name,
                "file_type": file_extension,
                "file_size": uploaded_file.size
            }
            
            # Process based on file type
            if file_extension == 'pdf':
                text = self._process_pdf(uploaded_file)
            elif file_extension in ['docx', 'doc']:
                text = self._process_docx(uploaded_file)
            elif file_extension == 'txt':
                text = self._process_txt(uploaded_file)
            elif file_extension == 'csv':
                text = self._process_csv(uploaded_file)
            elif file_extension in ['xlsx', 'xls']:
                text = self._process_excel(uploaded_file)
            elif file_extension in ['pptx', 'ppt']:
                text = self._process_pptx(uploaded_file)
            elif file_extension == 'json':
                text = self._process_json(uploaded_file)
            elif file_extension == 'html':
                text = self._process_html(uploaded_file)
            elif file_extension == 'xml':
                text = self._process_xml(uploaded_file)
            else:
                raise ValueError(f"Unsupported file format: {file_extension}")
            
            if not text.strip():
                raise ValueError("No text content extracted from file")
            
            # Split text into chunks
            chunks = self.text_splitter.split_text(text)
            
            # Create Document objects
            documents = []
            for i, chunk in enumerate(chunks):
                chunk_metadata = metadata.copy()
                chunk_metadata.update({
                    "chunk_id": i,
                    "total_chunks": len(chunks)
                })
                documents.append(Document(page_content=chunk, metadata=chunk_metadata))
            
            logger.info(f"Processed {uploaded_file.name}: {len(documents)} chunks created")
            return documents
            
        except Exception as e:
            logger.error(f"Error processing file {uploaded_file.name}: {str(e)}")
            raise e
    
    def _process_pdf(self, uploaded_file) -> str:
        """Process PDF file"""
        try:
            pdf_reader = PdfReader(uploaded_file)
            text = ""
            for page_num, page in enumerate(pdf_reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text += f"\n--- Page {page_num + 1} ---\n{page_text}\n"
            return text
        except Exception as e:
            logger.error(f"Error processing PDF: {str(e)}")
            raise e
    
    def _process_docx(self, uploaded_file) -> str:
        """Process DOCX file"""
        try:
            doc = DocxDocument(uploaded_file)
            text = ""
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text += paragraph.text + "\n"
            return text
        except Exception as e:
            logger.error(f"Error processing DOCX: {str(e)}")
            raise e
    
    def _process_txt(self, uploaded_file) -> str:
        """Process TXT file"""
        try:
            text = uploaded_file.read().decode('utf-8')
            return text
        except UnicodeDecodeError:
            # Try different encodings
            uploaded_file.seek(0)
            try:
                text = uploaded_file.read().decode('latin-1')
                return text
            except Exception as e:
                logger.error(f"Error decoding text file: {str(e)}")
                raise e
    
    def _process_csv(self, uploaded_file) -> str:
        """Process CSV file"""
        try:
            df = pd.read_csv(uploaded_file)
            text = f"CSV File: {uploaded_file.name}\n"
            text += f"Columns: {', '.join(df.columns.tolist())}\n"
            text += f"Number of rows: {len(df)}\n\n"
            
            # Convert to string representation
            text += "Data:\n"
            text += df.to_string(index=False)
            
            return text
        except Exception as e:
            logger.error(f"Error processing CSV: {str(e)}")
            raise e
    
    def _process_excel(self, uploaded_file) -> str:
        """Process Excel file"""
        try:
            # Read all sheets
            excel_file = pd.ExcelFile(uploaded_file)
            text = f"Excel File: {uploaded_file.name}\n"
            text += f"Sheets: {', '.join(excel_file.sheet_names)}\n\n"
            
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(uploaded_file, sheet_name=sheet_name)
                text += f"\n--- Sheet: {sheet_name} ---\n"
                text += f"Columns: {', '.join(df.columns.tolist())}\n"
                text += f"Number of rows: {len(df)}\n"
                text += df.to_string(index=False)
                text += "\n"
            
            return text
        except Exception as e:
            logger.error(f"Error processing Excel: {str(e)}")
            raise e
    
    def _process_pptx(self, uploaded_file) -> str:
        """Process PowerPoint file"""
        try:
            prs = Presentation(uploaded_file)
            text = f"PowerPoint File: {uploaded_file.name}\n"
            text += f"Total slides: {len(prs.slides)}\n\n"
            
            for slide_num, slide in enumerate(prs.slides, 1):
                text += f"\n--- Slide {slide_num} ---\n"
                
                # Extract text from shapes
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text += shape.text + "\n"
                    
                    # Extract text from tables
                    if shape.has_table:
                        table = shape.table
                        for row in table.rows:
                            row_text = []
                            for cell in row.cells:
                                row_text.append(cell.text)
                            text += " | ".join(row_text) + "\n"
            
            return text
        except Exception as e:
            logger.error(f"Error processing PPTX: {str(e)}")
            raise e
    
    def _process_json(self, uploaded_file) -> str:
        """Process JSON file"""
        try:
            json_data = json.load(uploaded_file)
            text = f"JSON File: {uploaded_file.name}\n\n"
            text += json.dumps(json_data, indent=2, ensure_ascii=False)
            return text
        except Exception as e:
            logger.error(f"Error processing JSON: {str(e)}")
            raise e
    
    def _process_html(self, uploaded_file) -> str:
        """Process HTML file"""
        try:
            html_content = uploaded_file.read().decode('utf-8')
            soup = BeautifulSoup(html_content, 'html.parser')
            
            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()
            
            # Get text content
            text = soup.get_text()
            
            # Clean up whitespace
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = '\n'.join(chunk for chunk in chunks if chunk)
            
            return f"HTML File: {uploaded_file.name}\n\n{text}"
        except Exception as e:
            logger.error(f"Error processing HTML: {str(e)}")
            raise e
    
    def _process_xml(self, uploaded_file) -> str:
        """Process XML file"""
        try:
            xml_content = uploaded_file.read().decode('utf-8')
            soup = BeautifulSoup(xml_content, 'xml')
            
            # Get text content
            text = soup.get_text()
            
            # Clean up whitespace
            lines = (line.strip() for line in text.splitlines())
            text = '\n'.join(line for line in lines if line)
            
            return f"XML File: {uploaded_file.name}\n\n{text}"
        except Exception as e:
            logger.error(f"Error processing XML: {str(e)}")
            raise e
    
    def get_document_summary(self, documents: List[Document]) -> Dict[str, Any]:
        """Get summary statistics for processed documents"""
        if not documents:
            return {}
        
        total_chars = sum(len(doc.page_content) for doc in documents)
        total_words = sum(len(doc.page_content.split()) for doc in documents)
        
        # Get unique files
        unique_files = set(doc.metadata.get('filename', 'Unknown') for doc in documents)
        
        return {
            "total_documents": len(documents),
            "total_characters": total_chars,
            "total_words": total_words,
            "unique_files": len(unique_files),
            "average_chunk_size": total_chars // len(documents) if documents else 0,
            "files": list(unique_files)
        }
